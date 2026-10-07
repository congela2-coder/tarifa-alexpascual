import json,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
d=json.load(open('data.json'))
F='Arial';wb=Workbook();ws=wb.active;ws.title='Tarifa'
red=PatternFill('solid',fgColor='FF2D2D');yel=PatternFill('solid',fgColor='FFF2A8');thin=Side(style='thin',color='D9D9D9')
ws['A1']='Precios del';ws['A1'].font=Font(name=F,bold=True)
ws['B1']=datetime.date(2026,10,7);ws['B1'].number_format='dd/mm/yyyy';ws['B1'].fill=yel;ws['B1'].font=Font(name=F,color='0000FF',bold=True)
ws['C1']='← Cambia la fecha cada viernes';ws['C1'].font=Font(name=F,italic=True,color='808080')
hd=['Código','Producto','Precio (€)','Familia']
for i,h in enumerate(hd,1):
    c=ws.cell(3,i,h);c.font=Font(name=F,bold=True,color='FFFFFF');c.fill=red;c.alignment=Alignment(horizontal='center')
fams=sorted({r[3] for r in d})
for r,(c,desc,p,f) in enumerate(d,4):
    for col,v in enumerate([c,desc,p,f],1):
        cell=ws.cell(r,col,v);cell.font=Font(name=F,color='0000FF');cell.border=Border(bottom=thin)
    ws.cell(r,1).number_format='@';ws.cell(r,3).number_format='#,##0.00 €'
for col,w in zip('ABCD',[12,48,13,26]):ws.column_dimensions[col].width=w
ws.freeze_panes='A4';ws.auto_filter.ref=f'A3:D{3+len(d)}'
fs=wb.create_sheet('Familias');fs['A1']='Familias';fs['A1'].font=Font(name=F,bold=True)
for i,f in enumerate(fams,2):fs.cell(i,1,f).font=Font(name=F)
fs.column_dimensions['A'].width=28
dv=DataValidation(type='list',formula1=f'=Familias!$A$2:$A$60',allow_blank=True,showErrorMessage=False)
dv.add('D4:D500');ws.add_data_validation(dv)
ins=wb.create_sheet('Instrucciones',0);ins.column_dimensions['A'].width=95
lines=[('Cómo actualizar la tarifa Alex Pascual',True),('',False),
('1. En la hoja «Tarifa», cambia la fecha amarilla de arriba (Precios del).',False),
('2. Cambia los precios en la columna «Precio (€)». Escribe solo el número: 11,24',False),
('3. Para añadir un producto, escríbelo en la primera fila vacía de abajo, con su código, nombre, precio y familia.',False),
('4. Para quitar un producto, borra su fila entera (clic derecho en el número de fila → Eliminar).',False),
('5. La familia se elige de la lista. Si necesitas una familia nueva, añádela en la hoja «Familias».',False),
('6. No cambies los títulos de las columnas ni el orden de las hojas.',False),
('7. Guarda el archivo y súbeselo a Claude escribiendo «actualiza la tarifa de Alex Pascual».',False),
('',False),('Los textos en azul son los datos que puedes cambiar. El código debe coincidir con el nombre de la foto en GitHub (ej. E5004 → fotos/E5004.jpg).',False)]
for i,(t,b) in enumerate(lines,1):
    c=ins.cell(i,1,t);c.font=Font(name=F,bold=b,size=14 if b else 11);c.alignment=Alignment(wrap_text=True)
wb.active=1
wb.save('/home/claude/out/Tarifa_Alex_Pascual.xlsx')
