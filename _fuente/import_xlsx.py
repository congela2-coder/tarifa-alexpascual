# Lee el Excel que manda Alex y regenera data.json + fecha
import json,sys,datetime
from openpyxl import load_workbook
ws=load_workbook(sys.argv[1],data_only=True)['Tarifa']
f=ws['B1'].value;fecha=f.strftime('%d/%m/%Y') if isinstance(f,(datetime.date,datetime.datetime)) else str(f)
rows=[]
for c,d,p,fam in ws.iter_rows(min_row=4,max_col=4,values_only=True):
    if not c:continue
    rows.append([str(c).strip(),str(d).strip(),round(float(str(p).replace('€','').replace(',','.').strip()),2),str(fam or '').strip().upper()])
json.dump(rows,open('data.json','w'),ensure_ascii=False);print(len(rows),'artículos',fecha)
